"""Build editable, unsigned product files from the catalog editorial manifest."""

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
COPY = json.loads((ROOT / "docs/catalog/five-product-editorial.json").read_text())
OUT = ROOT / "docs/partners/product-packets"
OUT.mkdir(parents=True, exist_ok=True)
PUBLIC = ROOT / "public/templates/parteneri"
PUBLIC.mkdir(parents=True, exist_ok=True)

TITLES = {
    "hp-keeper": "Cutiuță muzicală Harry Potter cu manivelă",
    "got-winter": "Cutiuță muzicală Game of Thrones Winter Is Coming",
    "kitten": "Cutiuță muzicală cu pisicuță și lună",
    "halloween": "Cutiuță muzicală Halloween cu manivelă",
    "sunshine": "Cutiuță muzicală You Are My Sunshine",
}


def table(doc, rows):
    tab = doc.add_table(rows=0, cols=2)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.style = "Light Shading Accent 1"
    for label, value in rows:
        cells = tab.add_row().cells
        cells[0].text = label
        cells[1].text = value
    doc.add_paragraph()


def heading(doc, title, text=None):
    doc.add_heading(title, level=1)
    if text:
        doc.add_paragraph(text)


def build(slug=None):
    item = COPY.get(slug, {})
    title = TITLES.get(slug, "Model universal pentru cutiuță muzicală")
    in_house = slug in {"kitten", "halloween"}
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor(35, 44, 54)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(21)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(20, 25, 31)
    title_props = styles["Title"]._element.pPr
    if title_props is not None:
        for border in title_props.findall(qn("w:pBdr")):
            title_props.remove(border)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 1"].font.color.rgb = RGBColor(41, 67, 72)
    footer = section.footer.paragraphs[0]
    footer.text = "Cutiuța Magică | Dosar produs editabil | Versiunea 1"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(100, 111, 119)

    doc.add_paragraph("Dosar tehnic și comercial pentru produs", style="Title")
    doc.add_paragraph(
        "Acest formular se completează pentru modelul și varianta de mai jos. "
        "Furnizorul sau producătorul anexează dovezile indicate, revizuiește afirmațiile comerciale "
        "și semnează numai informațiile pe care le poate confirma. O versiune nouă se încarcă în fișa produsului după fiecare actualizare."
    )
    table(doc, [
        ("Produs", title),
        ("SKU intern", item.get("sku", "[SKU]")),
        ("GTIN EAN oficial", "[De completat după alocarea GS1; nu folosiți un număr de test]"),
        ("Rol operator", "Producție proprie, de documentat" if in_house else "[Producător / furnizor / importator]"),
        ("Versiune, dată, lot", "[Versiune]  [Data]  [Lot sau serie]"),
    ])
    heading(doc, "1 Specificații și trasabilitate")
    table(doc, [
        ("Identitatea producătorului", "[Denumire, adresă poștală, e-mail, persoană de contact]"),
        ("Operator responsabil în UE", "[Denumire, adresă poștală și electronică, dacă este aplicabil]"),
        ("Materiale și componente", "[Lemn, finisaje, mecanism, accesorii, furnizori și coduri]"),
        ("Dimensiuni și greutate", "[Măsurători confirmate, toleranțe și metoda de măsurare]"),
        ("Funcționare și conținut pachet", "[Mecanism, melodie/variantă, ambalaj, elemente incluse]"),
        ("Origine și lanț de aprovizionare", "[Țară, fabrică/atelier, lot, factură și documente de origine]"),
        ("Fotografii și etichetă", "[Numele fișierelor aprobate și versiunea etichetei]"),
    ])

    doc.add_page_break()
    heading(doc, "2 Siguranță și conformitate", "Completați clasificarea produsului și măsurile aplicabile înainte de aprobarea fișei comerciale.")
    table(doc, [
        ("Destinație și grupa de vârstă", "[Obiect decorativ / obiect de colecție / jucărie; justificare și vârsta]"),
        ("Evaluare de risc", "[Riscuri mecanice, piese mici, muchii, finisaje, utilizare previzibilă]"),
        ("Teste și rapoarte", "[Standarde aplicabile, laborator, număr raport, dată și rezultat]"),
        ("Marcaje aplicabile", "[Identificare model/lot, producător, avertismente, CE numai dacă este necesar]"),
        ("Instrucțiuni pentru client", "[Utilizare, întreținere, depozitare, avertismente în română]"),
        ("Documente anexate", "[Evaluare de risc, rapoarte, declarații, macheta etichetei]"),
    ])
    doc.add_paragraph(
        "Confirmare de completat: datele și documentele anexate corespund variantei și lotului indicate. "
        "Clasificarea și marcajele au fost stabilite în raport cu cerințele aplicabile produsului."
    )
    table(doc, [("Responsabil și funcție", "[Nume și funcție]"), ("Semnătură și dată", "[Semnătură]  [Data]")])

    doc.add_page_break()
    heading(doc, "3 Drepturi și autorizare comercială")
    doc.add_paragraph(
        "Completați separat pentru fiecare element protejat. O factură de achiziție a produsului nu ține locul "
        "unei licențe pentru marcă, ilustrație, compoziție muzicală, înregistrare sau material publicitar."
    )
    table(doc, [
        ("Marcă, nume și elemente tematice", "[Titular, lanț de autorizare, document anexat]"),
        ("Ilustrație și design", "[Autor/titular, cesiune ori licență, versiune și dovezi]"),
        ("Melodie și înregistrare", "[Titlu, compozitor, titulari, mecanism, licență și utilizări permise]"),
        ("Fotografii, video și texte", "[Creator, titular, consimțământ sau licență, versiune]"),
        ("Canale, teritoriu și perioadă", "[Magazin, marketplace, social, publicitate; țări; început și sfârșit]"),
        ("Drept de revânzare și promovare", "[Produse/loturi acoperite, limitări, sub-licențiere dacă există]"),
        ("Anexe", "[Contract, licență, factură, acord de utilizare, alte dovezi]"),
    ])
    doc.add_paragraph(
        "Autorizare de completat de titularul drepturilor sau reprezentantul împuternicit: "
        "confirm că drepturile și utilizările descrise mai sus sunt acoperite de documentele anexate."
    )
    table(doc, [("Entitate și calitate", "[Denumire și calitate]"), ("Semnătură și dată", "[Semnătură]  [Data]")])

    doc.add_page_break()
    heading(doc, "4 Controlul lotului și aprobarea fișei")
    if in_house:
        doc.add_paragraph(
            "Pentru acest model, fișa este destinată producției proprii. Atelierul completează rețeta materialelor, "
            "versiunile desenelor și verificările pe lot înainte de prima livrare."
        )
    table(doc, [
        ("Lot și cantitate", "[Cod lot, data producției/recepției, unități]"),
        ("Componente și furnizori", "[Coduri componente, loturi de intrare, furnizori]"),
        ("Pași de asamblare", "[Operații, versiunea instrucțiunilor, responsabil]"),
        ("Verificări finale", "[Aspect, fixare, manivelă, funcționare, etichetă, ambalaj]"),
        ("Neconformități și remediere", "[Descriere, unități afectate, acțiune și aprobare]"),
        ("Documente predate partenerilor", "[Fișă tehnică, instrucțiuni, fotografii, acte și versiuni]"),
    ])
    heading(doc, "Confirmarea informațiilor pentru listare")
    doc.add_paragraph("Titlu propus: " + item.get("marketplaceTitle", "[Titlu]"))
    doc.add_paragraph("Descriere propusă: " + item.get("marketplaceDescription", "[Descriere]"))
    doc.add_paragraph(
        "Verificatorul marchează corecturile necesare în document, iar administratorul actualizează "
        "fișa produsului după primirea dovezilor."
    )
    table(doc, [
        ("Persoana care confirmă datele", "[Nume, rol, semnătură, dată]"),
        ("Administrator care verifică dosarul", "[Nume, rol, semnătură, dată]"),
    ])
    name = slug or "model-universal"
    target = OUT / f"{name}-dosar-produs.docx"
    doc.save(target)
    shutil.copyfile(target, PUBLIC / target.name)
    return target


if __name__ == "__main__":
    for product_slug in [None, *COPY]:
        print(build(product_slug))
